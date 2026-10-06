#!/usr/bin/env python3
"""Xuất file Excel tổng hợp bug UI.

Usage:
    python export_bugs.py bugs.json output.xlsx [--project "SecondPhone Web"] [--tester "HoaLT"]

bugs.json: list các object, giữ đúng thứ tự màn trong flow:
    [{"screen": "Intro 1",
      "title": "[UI][Intro 1] Title bị lệch lên trên so với Design",
      "severity": "Medium",
      "actual": "- ...", "expected": "- ...",
      "attachment": "Screenshot/Video: "}]
Màn không có lỗi: {"screen": "Intro 2", "no_bug": true}
"""
import argparse, json, re, datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

HEAD_FILL = PatternFill("solid", start_color="1F4E78")
HEAD_FONT = Font(name="Arial", bold=True, color="FFFFFF")
BODY_FONT = Font(name="Arial", size=10)
SCREEN_FILL = PatternFill("solid", start_color="DDEBF7")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
SEVERITIES = ["Critical", "High", "Medium", "Low"]
SEV_FILL = {k: PatternFill("solid", start_color=c) for k, c in
            zip(SEVERITIES, ["F4B6B6", "F8CBAD", "FFE699", "E2EFDA"])}

def bug_type(title):
    tags = re.findall(r"\[([^\]]+)\]", title or "")
    if not tags:
        return ""
    # [UI][Responsive][Intro 1] -> [UI][Responsive] ; [UI/UX][Paywall 1] -> [UI/UX]
    known = {"UI", "UI/UX", "Logic", "API", "Crash", "Performance", "Responsive"}
    types = [t for t in tags if t in known]
    return "".join(f"[{t}]" for t in types) if types else f"[{tags[0]}]"

def style_header(ws, row, ncol):
    for c in range(1, ncol + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill, cell.font, cell.alignment, cell.border = HEAD_FILL, HEAD_FONT, CENTER, BORDER

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("bugs_json"); ap.add_argument("output")
    ap.add_argument("--project", default=""); ap.add_argument("--tester", default="")
    a = ap.parse_args()
    items = json.load(open(a.bugs_json, encoding="utf-8"))

    wb = Workbook()
    ws = wb.active; ws.title = "Bug List"
    headers = ["#", "Screen", "Bug Type", "Severity", "Title", "Actual Result", "Expected Result", "Attachment"]
    widths = [5, 16, 13, 10, 45, 50, 50, 26]
    ws.append(headers); style_header(ws, 1, len(headers))
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    screens, n, r = [], 0, 2
    for it in items:
        sc = it.get("screen", "").strip()
        if sc and sc not in screens:
            screens.append(sc)
        if it.get("no_bug"):
            vals = ["", sc, "", "", "Không phát hiện lỗi", "", "", ""]
        else:
            n += 1
            vals = [n, sc, bug_type(it.get("title", "")), it.get("severity", ""), it.get("title", ""),
                    it.get("actual", ""), it.get("expected", ""),
                    it.get("attachment", "Screenshot/Video: ")]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font, cell.border = BODY_FONT, BORDER
            cell.alignment = CENTER if c in (1, 2, 3, 4) else WRAP
            if c == 4 and v in SEV_FILL:
                cell.fill = SEV_FILL[v]
            if it.get("no_bug"):
                cell.font = Font(name="Arial", size=10, italic=True, color="808080")
        r += 1
    last = r - 1
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(last,1)}"

    # Summary sheet: đếm theo Screen x Bug Type bằng công thức
    sm = wb.create_sheet("Summary")
    sm["A1"] = "Tổng hợp bug"; sm["A1"].font = Font(name="Arial", bold=True, size=14)
    sm["A2"] = f"Project: {a.project}" if a.project else "Project:"
    sm["A3"] = f"Tester: {a.tester}" if a.tester else "Tester:"
    sm["A4"] = f"Ngày export: {datetime.date.today():%d/%m/%Y}"
    for c in ("A2", "A3", "A4"):
        sm[c].font = BODY_FONT
    types = sorted({bug_type(it.get("title", "")) for it in items if not it.get("no_bug")} - {""})
    hr = 6
    head = ["Screen"] + types + ["Tổng"]
    for c, v in enumerate(head, 1):
        sm.cell(row=hr, column=c, value=v)
    style_header(sm, hr, len(head))
    rng_s, rng_t = f"'Bug List'!$B$2:$B${last}", f"'Bug List'!$C$2:$C${last}"
    for i, sc in enumerate(screens):
        rr = hr + 1 + i
        sm.cell(row=rr, column=1, value=sc)
        for j, t in enumerate(types):
            col = get_column_letter(2 + j)
            sm.cell(row=rr, column=2 + j,
                    value=f'=COUNTIFS({rng_s},$A{rr},{rng_t},{col}${hr})')
        tc = 2 + len(types)
        sm.cell(row=rr, column=tc,
                value=f"=SUM({get_column_letter(2)}{rr}:{get_column_letter(tc-1)}{rr})" if types else 0)
        for c in range(1, len(head) + 1):
            cell = sm.cell(row=rr, column=c); cell.border, cell.font = BORDER, BODY_FONT
            cell.alignment = CENTER if c > 1 else Alignment(vertical="center")
    tr = hr + 1 + len(screens)
    sm.cell(row=tr, column=1, value="Tổng")
    for c in range(2, len(head) + 1):
        L = get_column_letter(c)
        sm.cell(row=tr, column=c, value=f"=SUM({L}{hr+1}:{L}{tr-1})" if screens else 0)
    for c in range(1, len(head) + 1):
        cell = sm.cell(row=tr, column=c)
        cell.font, cell.border, cell.fill = Font(name="Arial", bold=True), BORDER, SCREEN_FILL
        cell.alignment = CENTER if c > 1 else Alignment(vertical="center")
    # Severity counts
    sr = tr + 2
    for c, v in enumerate(["Severity", "Số bug"], 1):
        sm.cell(row=sr, column=c, value=v)
    style_header(sm, sr, 2)
    rng_v = f"'Bug List'!$D$2:$D${last}"
    for i, sv in enumerate(SEVERITIES):
        rr = sr + 1 + i
        sm.cell(row=rr, column=1, value=sv).fill = SEV_FILL[sv]
        sm.cell(row=rr, column=2, value=f"=COUNTIF({rng_v},A{rr})")
        for c in (1, 2):
            cell = sm.cell(row=rr, column=c); cell.border, cell.font = BORDER, BODY_FONT
    sm.column_dimensions["A"].width = 20
    for c in range(2, len(head) + 1):
        sm.column_dimensions[get_column_letter(c)].width = 13
    wb.move_sheet("Summary", offset=-1)
    wb.active = 0
    wb.save(a.output)
    print(f"Saved {a.output}: {n} bugs, {len(screens)} screens")

if __name__ == "__main__":
    main()
